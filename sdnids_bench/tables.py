"""Manuscript tables for Paper 2: CSV (data), LaTeX (booktabs) and one Word file with every table."""
from __future__ import annotations

import json
import os

import pandas as pd

from .experiments import SEARCH_SPACE_TEXT, ds_dir
from .models import FAMILY


def pm(mean, sd, d=4):
    if pd.isna(mean):
        return "–"
    return f"{mean:.{d}f} ± {sd:.{d}f}" if not pd.isna(sd) else f"{mean:.{d}f}"


def pval(p):
    if pd.isna(p):
        return "–"
    return "<0.001" if p < 0.001 else f"{p:.3f}"


# ------------------------------------------------------------------ Table 1
def table_dataset_profile(profiles: dict) -> pd.DataFrame:
    rows = []
    for ds, p in profiles.items():
        q = p.copy()
        q.insert(0, "Dataset", ds)
        rows.append(q)
    t = pd.concat(rows, ignore_index=True)
    for c in ("Source flows", "Loaded flows", "After deduplication"):
        t[c] = t[c].map(lambda v: f"{int(v):,}")
    for c in ("Share after dedup. (%)", "Removed by dedup. (%)"):
        t[c] = t[c].map(lambda v: f"{v:.2f}")
    return t


# ------------------------------------------------------------------ Table 2
def table_hyperparams(cfg, datasets) -> pd.DataFrame:
    rows = []
    for ds in datasets:
        p = os.path.join(ds_dir(cfg, ds), "best_hp.json")
        if not os.path.exists(p):
            continue
        best = json.load(open(p))
        for k, hp in best.items():
            model, task = k.split("|")
            rows.append({"Dataset": ds, "Model": model, "Family": FAMILY.get(model, ""), "Task": task,
                         "Search space": SEARCH_SPACE_TEXT["RF" if model == "RF" else "XGBoost" if model == "XGBoost"
                                                           else "neural"],
                         "Selected": "; ".join(f"{a}={b}" for a, b in hp.items())})
    return pd.DataFrame(rows)


# ------------------------------------------------------------------ Table 3
MAIN_METRICS = [("macro_f1", "Macro-F1"), ("macro_f1_min10", "Macro-F1 (classes ≥10 test flows)"), ("pr_auc_macro", "PR-AUC"), ("mcc", "MCC"),
                ("balanced_acc", "Bal. acc."), ("minority_recall", "Minority recall"),
                ("fpr", "FPR"), ("fnr", "FNR"), ("accuracy", "Accuracy")]
LOWER_BETTER = {"fpr", "fnr", "latency_ms_b1", "size_mb", "params", "train_time_s"}


def table_main_results(res: pd.DataFrame, boot: pd.DataFrame | None, dataset, task, L, exp="B1"):
    g = res[(res.exp == exp) & (res.dataset == dataset) & (res.task == task) & (res.L == L) & (res.protocol == "SC")]
    metrics = [(m, lab) for m, lab in MAIN_METRICS if m in g and g[m].notna().any()]
    agg = g.groupby("model")[[m for m, _ in metrics]].agg(["mean", "std"])
    order = agg[("macro_f1", "mean")].sort_values(ascending=False).index
    rows, best = [], {}
    for m, _ in metrics:
        col = agg[(m, "mean")]
        best[m] = col.idxmin() if m in LOWER_BETTER else col.idxmax()
    for model in order:
        r = {"Model": model, "Seeds": int(g[g.model == model]["seed"].nunique())}
        for m, lab in metrics:
            r[lab] = pm(agg.loc[model, (m, "mean")], agg.loc[model, (m, "std")])
            if best[m] == model:
                r[lab] = "**" + r[lab] + "**"
        if boot is not None and len(boot):
            b = boot[(boot.dataset == dataset) & (boot.task == task) & (boot.L == L) & (boot.model == model) &
                     (boot.protocol == "SC") & (boot.exp == exp)]
            if len(b):
                r["Macro-F1 95% CI"] = f"[{b.macro_f1_ci_low.iloc[0]:.4f}, {b.macro_f1_ci_high.iloc[0]:.4f}]"
        rows.append(r)
    return pd.DataFrame(rows)


# ------------------------------------------------------------------ Table 4
def table_context(res, dataset, metric="macro_f1"):
    g = res[(res.exp == "B3") & (res.dataset == dataset) & (res.task == "multi")]
    a = g.groupby(["model", "L"])[metric].agg(["mean", "std"]).reset_index()
    a["v"] = [pm(m, s) for m, s in zip(a["mean"], a["std"])]
    t = a.pivot(index="model", columns="L", values="v")
    t.columns = [f"L = {c}" for c in t.columns]
    return t.reset_index().rename(columns={"model": "Model"})


# ------------------------------------------------------------------ Table 6
def table_imbalance(res, dataset, classes):
    g = res[(res.exp == "B4") & (res.dataset == dataset) & (res.task == "multi")]
    cols = [("macro_f1", "Macro-F1"), ("minority_recall", "Minority recall"), ("fpr", "FPR")]
    minor_idx = [(i, c) for i, c in enumerate(classes) if c in ("BFA", "Web-Attack", "BOTNET", "U2R", "Other")]
    cols += [(f"recall_c{i}", f"Recall {c}") for i, c in minor_idx]
    a = g.groupby(["model", "imbalance"])[[c for c, _ in cols]].agg(["mean", "std"])
    rows = []
    for (model, imb), r in a.iterrows():
        label = {"none": "None", "weighted_ce": "Class weighting" if model in ("RF", "XGBoost") else "Weighted CE",
                 "focal": "Focal loss", "logit_adjusted": "Logit-adjusted loss",
                 "balanced_sampling": "Balanced sampling", "smote": "SMOTE"}.get(imb, imb)
        row = {"Model": model, "Strategy": label}
        for c, lab in cols:
            row[lab] = pm(r[(c, "mean")], r[(c, "std")], 3)
        rows.append(row)
    return pd.DataFrame(rows)


# ------------------------------------------------------------------ Table 7
def table_cost(res, dataset):
    g = res[(res.exp == "B5") & (res.dataset == dataset)]
    a = g.groupby("model")[["macro_f1", "params", "size_mb", "flops_per_flow", "latency_ms_b1",
                            "latency_ms_b1_p95", "throughput_fps", "train_time_s"]].mean()
    t = pd.DataFrame({
        "Model": a.index,
        "Macro-F1": a["macro_f1"].map(lambda v: f"{v:.4f}"),
        "Parameters / tree nodes": a["params"].map(lambda v: "–" if pd.isna(v) else f"{int(v):,}"),
        "Size (MB)": a["size_mb"].map(lambda v: f"{v:.2f}"),
        "FLOPs per window": a["flops_per_flow"].map(lambda v: "–" if pd.isna(v) else f"{int(v):,}"),
        "Latency b=1 (ms, median)": a["latency_ms_b1"].map(lambda v: f"{v:.3f}"),
        "Latency b=1 (ms, p95)": a["latency_ms_b1_p95"].map(lambda v: f"{v:.3f}"),
        "Throughput (windows/s)": a["throughput_fps"].map(lambda v: f"{v:,.0f}"),
        "Training time (s)": a["train_time_s"].map(lambda v: f"{v:,.0f}"),
    })
    return t.sort_values("Macro-F1", ascending=False).reset_index(drop=True)


# ------------------------------------------------------------------ export
def _latex(df: pd.DataFrame, caption, label):
    body = df.copy()
    for c in body.columns:
        body[c] = body[c].astype(str).str.replace(r"\*\*(.*)\*\*", r"\\textbf{\1}", regex=True) \
            .str.replace("±", r"$\pm$", regex=False).str.replace("%", r"\%", regex=False) \
            .str.replace("_", r"\_", regex=False).str.replace("Δ", r"$\Delta$", regex=False) \
            .str.replace("τ", r"$\tau$", regex=False).str.replace("≥", r"$\geq$", regex=False)
    cols = [str(c).replace("%", r"\%").replace("_", r"\_").replace("Δ", r"$\Delta$").replace("τ", r"$\tau$").replace("≥", r"$\geq$")
            for c in body.columns]
    body.columns = cols
    caption = caption.replace("τ", r"$\tau$").replace("%", r"\%")
    tex = body.to_latex(index=False, escape=False, column_format="l" * len(cols))
    tex = tex.replace("\\toprule", "\\toprule").replace("\\begin{tabular}", "\\resizebox{\\textwidth}{!}{%\n\\begin{tabular}") \
        .replace("\\end{tabular}", "\\end{tabular}}")
    return (f"\\begin{{table*}}[t]\n\\centering\n\\caption{{{caption}}}\n\\label{{{label}}}\n{tex}\\end{{table*}}\n")


def export_tables(tables: dict, out_dir: str, docx_name="Paper2_Tables.docx"):
    """tables: ordered {short_name: (DataFrame, caption)}. Writes CSV per table, tables.tex and a .docx."""
    os.makedirs(out_dir, exist_ok=True)
    tex_parts = ["% Requires \\usepackage{booktabs,graphicx}\n"]
    for name, (df, cap) in tables.items():
        df.to_csv(os.path.join(out_dir, f"{name}.csv"), index=False)
        tex_parts.append(_latex(df, cap, f"tab:{name}"))
    with open(os.path.join(out_dir, "tables.tex"), "w", encoding="utf-8") as f:
        f.write("\n".join(tex_parts))
    _docx(tables, os.path.join(out_dir, docx_name))
    return out_dir


def _docx(tables, path):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Inches

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.7))
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(10)
    doc.add_heading("Paper 2 — Manuscript tables", level=1)

    def border(cell, **edges):
        tcPr = cell._tc.get_or_add_tcPr()
        b = tcPr.find(qn("w:tcBorders"))
        if b is None:
            b = OxmlElement("w:tcBorders")
            tcPr.append(b)
        for edge, sz in edges.items():
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:color"), "000000")
            b.append(el)

    for i, (name, (df, cap)) in enumerate(tables.items(), 1):
        p = doc.add_paragraph()
        r = p.add_run(f"Table {i}. ")
        r.bold = True
        p.add_run(cap)
        t = doc.add_table(rows=1, cols=len(df.columns))
        t.autofit = True
        for j, c in enumerate(df.columns):
            cell = t.rows[0].cells[j]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(c))
            run.bold, run.font.size = True, Pt(8)
            border(cell, top=12, bottom=6)
        for _, row in df.iterrows():
            cells = t.add_row().cells
            for j, v in enumerate(row):
                s = "" if pd.isna(v) else str(v)
                bold = s.startswith("**") and s.endswith("**")
                run = cells[j].paragraphs[0].add_run(s.strip("*"))
                run.bold, run.font.size = bold, Pt(8)
        for cell in t.rows[-1].cells:
            border(cell, bottom=12)
        doc.add_paragraph()
    doc.save(path)
